"""Read documents and split them into passages ("chunks") that remember their page, heading and position."""
from __future__ import annotations

import json
import logging
import re
import statistics
import threading
from dataclasses import dataclass, field
from pathlib import Path

from . import db, search

log = logging.getLogger("askdocs.ingest")
_pdfium_lock = threading.Lock()

SUPPORTED = {".pdf", ".docx", ".txt", ".md"}
TARGET_CHARS = 850


class IngestError(Exception):
    pass


@dataclass
class Para:
    text: str
    heading: bool = False
    page: int | None = None
    boxes: list = field(default_factory=list)
    page_w: float | None = None
    page_h: float | None = None


# --------------------------------------------------------------------------- readers
def _pdf_paras(path: Path) -> tuple[list[Para], int]:
    import pdfplumber
    paras: list[Para] = []
    try:
        pdf = pdfplumber.open(str(path))
    except Exception:
        raise IngestError("This PDF seems damaged or password-protected. Please export it again and re-upload.")
    with pdf:
        n_pages = len(pdf.pages)
        all_sizes = []
        pages_lines = []
        for pno, page in enumerate(pdf.pages, start=1):
            words = page.extract_words(extra_attrs=["size", "fontname"], use_text_flow=True)
            lines = []
            for w in sorted(words, key=lambda w: (round(w["top"] / 3), w["x0"])):
                if lines and abs(lines[-1]["top"] - w["top"]) < 3:
                    ln = lines[-1]
                    ln["words"].append(w)
                    ln["x1"] = max(ln["x1"], w["x1"]); ln["bottom"] = max(ln["bottom"], w["bottom"])
                else:
                    lines.append({"words": [w], "top": w["top"], "bottom": w["bottom"], "x0": w["x0"], "x1": w["x1"]})
            clean = []
            for ln in lines:
                ln["words"].sort(key=lambda w: w["x0"])
                ln["text"] = " ".join(w["text"] for w in ln["words"])
                ln["size"] = statistics.median([w["size"] for w in ln["words"]])
                ln["bold"] = all("Bold" in w.get("fontname", "") for w in ln["words"])
                if ln["top"] > page.height * 0.93:  # running footer
                    continue
                clean.append(ln)
                all_sizes += [ln["size"]] * len(ln["text"])
            pages_lines.append((pno, page.width, page.height, clean))
        if not all_sizes:
            raise IngestError("This PDF has no readable text (it may be a scan). Please upload a text PDF or Word file.")
        body = statistics.median(all_sizes)
        for pno, pw, ph, lines in pages_lines:
            cur = None
            for ln in lines:
                is_head = ln["size"] >= body + 1.5 or (ln["bold"] and len(ln["text"]) < 80 and not ln["text"].endswith("."))
                gap = ln["top"] - cur["bottom"] if cur else 999
                new_para = cur is None or is_head != cur["heading"] or gap > ln["size"] * 0.9 or is_head
                if new_para:
                    if cur:
                        paras.append(_finish(cur, pno, pw, ph))
                    cur = {"lines": [ln], "heading": is_head, "bottom": ln["bottom"]}
                else:
                    cur["lines"].append(ln); cur["bottom"] = ln["bottom"]
            if cur:
                paras.append(_finish(cur, pno, pw, ph))
    return paras, n_pages


def _finish(cur, pno, pw, ph) -> Para:
    lines = cur["lines"]

    def is_row(l):
        # table rows have wide gaps between columns; prose lines only have normal word spaces
        ws = l["words"]
        wide = sum(1 for a, b in zip(ws, ws[1:]) if b["x0"] - a["x1"] > l["size"] * 1.2)
        return wide >= 2
    tableish = len(lines) >= 1 and sum(is_row(l) for l in lines) >= max(1, 0.6 * len(lines))
    sep = "\n" if tableish else " "
    text = sep.join(l["text"] for l in lines)
    text = re.sub(r"(\w)- (\w)", r"\1\2", text) if not tableish else text
    boxes = [[round(l["x0"], 1), round(l["top"], 1), round(l["x1"], 1), round(l["bottom"], 1)] for l in lines]
    return Para(text, cur["heading"], pno, boxes, pw, ph)


def _docx_paras(path: Path) -> tuple[list[Para], int]:
    from docx import Document
    try:
        d = Document(str(path))
    except Exception:
        raise IngestError("We couldn't open this Word file. Please save it again as .docx and re-upload.")
    paras = []
    for p in d.paragraphs:
        t = p.text.strip()
        if not t:
            continue
        style = (p.style.name or "").lower()
        paras.append(Para(t, heading=style.startswith("heading") or style == "title"))
    for table in d.tables:
        rows = [" | ".join(c.text.strip() for c in r.cells) for r in table.rows]
        if rows:
            paras.append(Para("\n".join(rows)))
    return paras, 0


def _txt_paras(path: Path) -> tuple[list[Para], int]:
    raw = path.read_text(encoding="utf-8", errors="replace")
    paras = []
    for block in re.split(r"\n\s*\n", raw):
        lines = [l.strip() for l in block.strip().splitlines() if l.strip()]
        if not lines:
            continue
        if lines[0].startswith("#"):
            paras.append(Para(lines[0].lstrip("# "), heading=True))
            lines = lines[1:]
        elif len(lines) >= 2 and lines[0].endswith("?"):
            # FAQ style: the question is the heading of its answer
            paras.append(Para(lines[0], heading=True))
            lines = lines[1:]
        if lines:
            paras.append(Para(" ".join(lines)))
    return paras, 0


def read_paras(path: Path):
    ext = path.suffix.lower()
    if ext == ".pdf":
        return _pdf_paras(path)
    if ext == ".docx":
        return _docx_paras(path)
    if ext in (".txt", ".md"):
        return _txt_paras(path)
    raise IngestError("Please upload PDF, Word (.docx) or text files.")


# --------------------------------------------------------------------------- chunking
def make_chunks(paras: list[Para], title: str) -> list[dict]:
    chunks, heading, cur = [], "", None
    first_heading = next((p.text.strip() for p in paras if p.heading), None)

    def flush():
        nonlocal cur
        if cur and cur["text"].strip():
            # skip the "Company · Version 2, effective ..." line under the document title
            if not (cur["heading"] == first_heading and len(cur["text"]) < 120 and re.search(r"version|updated|effective|valid|·", cur["text"], re.I)):
                chunks.append(cur)
        cur = None

    for p in paras:
        if p.heading:
            flush()
            heading = p.text.strip()
            continue
        if cur and (cur["page"] != p.page or len(cur["text"]) + len(p.text) > TARGET_CHARS):
            flush()
        if cur is None:
            cur = {"page": p.page, "heading": heading, "text": "", "boxes": [], "page_w": p.page_w, "page_h": p.page_h}
        cur["text"] = (cur["text"] + "\n" + p.text).strip()
        cur["boxes"] += p.boxes
    flush()
    return chunks


def ingest_document(doc_id: int):
    d = db.one("SELECT * FROM documents WHERE id=?", (doc_id,))
    try:
        db.execute("UPDATE documents SET status='processing', error=NULL WHERE id=?", (doc_id,))
        paras, n_pages = read_paras(Path(d["stored_path"]))
        title = d["title"] or Path(d["filename"]).stem.replace("_", " ")
        chunks = make_chunks(paras, title)
        if not chunks:
            raise IngestError("We couldn't find any text in this file.")
        vectors = search.embed_passages([f"{title}. {c['heading']}. {c['text']}" for c in chunks])
        with db.tx() as c:
            c.execute("DELETE FROM chunks WHERE doc_id=?", (doc_id,))
            for i, ch in enumerate(chunks):
                c.execute("INSERT INTO chunks(doc_id, kb_id, position, page, heading, text, boxes_json, page_w, page_h, embedding) VALUES(?,?,?,?,?,?,?,?,?,?)",
                          (doc_id, d["kb_id"], i, ch["page"], ch["heading"], ch["text"], json.dumps(ch["boxes"]), ch["page_w"], ch["page_h"],
                           vectors[i].tobytes() if vectors is not None else None))
            c.execute("UPDATE documents SET status='ready', pages=?, chunks=?, updated_at=? WHERE id=?", (n_pages, len(chunks), db.now(), doc_id))
    except IngestError as e:
        db.execute("UPDATE documents SET status='failed', error=? WHERE id=?", (str(e), doc_id))
    except Exception as e:  # pragma: no cover
        log.exception("ingest failed")
        db.execute("UPDATE documents SET status='failed', error=? WHERE id=?", ("Something went wrong reading this file. Please try again.", doc_id))
    search.invalidate(d["kb_id"])


def ingest_async(doc_id: int):
    threading.Thread(target=ingest_document, args=(doc_id,), daemon=True).start()


def render_page(doc_id: int, page: int) -> Path | None:
    """PNG of one PDF page for the document viewer (cached)."""
    from .config import PAGES_DIR
    out = PAGES_DIR / f"{doc_id}_{page}.png"
    if out.exists():
        return out
    d = db.one("SELECT stored_path, file_type FROM documents WHERE id=?", (doc_id,))
    if not d or d["file_type"] != "pdf" or not d["stored_path"] or not Path(d["stored_path"]).exists():
        return None
    import pypdfium2 as pdfium
    with _pdfium_lock:
        pdf = pdfium.PdfDocument(d["stored_path"])
        try:
            if page < 1 or page > len(pdf):
                return None
            pdf[page - 1].render(scale=2).to_pil().save(out)
        finally:
            pdf.close()
    return out
