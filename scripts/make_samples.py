"""Generate the demo company's documents (Nimbus Home Appliances, a fictitious company).

    python -m scripts.make_samples

Writes sample_data/docs/*.pdf|docx|txt. Every document says it is sample content.
"""
from __future__ import annotations

from pathlib import Path

from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

OUT = Path(__file__).resolve().parent.parent / "sample_data" / "docs"
COMPANY = "Nimbus Home Appliances (Pvt) Ltd"
SAMPLE_NOTE = "SAMPLE DOCUMENT - Nimbus Home Appliances is a fictitious company created for demonstration."

ss = getSampleStyleSheet()
H1 = ParagraphStyle("h1", parent=ss["Heading1"], fontName="Helvetica-Bold", fontSize=20, textColor=colors.HexColor("#0f766e"), spaceAfter=6)
H2 = ParagraphStyle("h2", parent=ss["Heading2"], fontName="Helvetica-Bold", fontSize=13.5, textColor=colors.HexColor("#111827"), spaceBefore=10, spaceAfter=4)
P = ParagraphStyle("p", parent=ss["BodyText"], fontName="Helvetica", fontSize=10.5, leading=15, alignment=TA_LEFT, spaceAfter=5)
META = ParagraphStyle("meta", parent=P, fontSize=9, textColor=colors.HexColor("#6b7280"))


def _footer(title):
    def draw(c, doc):
        c.saveState()
        c.setFont("Helvetica", 8)
        c.setFillColor(colors.HexColor("#6b7280"))
        c.drawString(20 * mm, 12 * mm, f"{COMPANY} · {title}")
        c.drawRightString(190 * mm, 12 * mm, f"Page {doc.page}")
        c.drawString(20 * mm, 8 * mm, SAMPLE_NOTE)
        c.restoreState()
    return draw


def build_pdf(name, title, version, blocks):
    doc = SimpleDocTemplate(str(OUT / name), pagesize=A4, leftMargin=20 * mm, rightMargin=20 * mm, topMargin=18 * mm, bottomMargin=22 * mm,
                            title=title, author=COMPANY)
    story = [Paragraph(title, H1), Paragraph(f"{COMPANY} · {version}", META), Spacer(1, 6)]
    for b in blocks:
        if b == "PAGEBREAK":
            story.append(PageBreak())
        elif isinstance(b, tuple) and b[0] == "h":
            story.append(Paragraph(b[1], H2))
        elif isinstance(b, tuple) and b[0] == "table":
            t = Table(b[1], colWidths=b[2] if len(b) > 2 else None, repeatRows=1)
            t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0f766e")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                                   ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
                                   ("FONTSIZE", (0, 0), (-1, -1), 9.5), ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#d1d5db")),
                                   ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f3f4f6")]),
                                   ("VALIGN", (0, 0), (-1, -1), "TOP"), ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4)]))
            story += [t, Spacer(1, 6)]
        else:
            story.append(Paragraph(b, P))
    doc.build(story, onFirstPage=_footer(title), onLaterPages=_footer(title))


def hr_policy():
    build_pdf("HR_Policy.pdf", "Human Resources Policy", "Version 3.2, effective 1 January 2026", [
        ("h", "1. Purpose and scope"),
        "This policy explains working hours, leave, benefits and conduct rules for all employees of Nimbus Home Appliances. It applies to "
        "permanent, contract and part-time staff at the head office, warehouses and retail branches unless a section says otherwise.",
        ("h", "2. Working hours"),
        "Office hours are 9:00 am to 6:00 pm, Monday to Friday, with a one-hour lunch break from 1:00 pm to 2:00 pm. Retail branch staff work "
        "shifts published by the branch manager each week, with one weekly day off. During Ramadan, office hours are 9:00 am to 3:00 pm.",
        "Employees who arrive more than 15 minutes late three times in a calendar month will have half a day deducted from casual leave.",
        ("h", "3. Probation and notice period"),
        "New employees serve a probation period of 3 months. During probation, either side may end employment with 1 week's notice. "
        "After confirmation, the notice period is 1 month for staff and 2 months for managers.",
        ("h", "4. Leave entitlements"),
        "Employees are entitled to 10 casual leaves per calendar year. Casual leave cannot be carried forward to the next year and cannot be "
        "encashed. Casual leave of more than 2 consecutive days needs approval from the line manager at least 3 days in advance.",
        "Employees are entitled to 8 sick leaves per calendar year. Sick leave of more than 2 consecutive days requires a medical certificate.",
        "Annual leave is 14 working days per year, earned after confirmation. Up to 7 unused annual leave days may be carried forward to the next "
        "year; any balance above 7 days lapses on 31 December.",
        "Part-time staff receive leave on a pro-rata basis: 5 casual leaves and 4 sick leaves per calendar year, and 7 days of annual leave.",
        ("table", [["Leave type", "Full-time", "Part-time", "Carry forward"],
                   ["Casual leave", "10 days", "5 days", "No"],
                   ["Sick leave", "8 days", "4 days", "No"],
                   ["Annual leave", "14 days", "7 days", "Up to 7 days"],
                   ["Maternity leave", "90 days paid", "90 days paid", "-"],
                   ["Paternity leave", "7 days paid", "7 days paid", "-"]], [45 * mm, 35 * mm, 35 * mm, 40 * mm]),
        "Maternity leave is 90 days fully paid and can be taken up to 30 days before the expected delivery date. Paternity leave is 7 days fully "
        "paid, to be taken within one month of the child's birth.",
        "PAGEBREAK",
        ("h", "5. Public holidays"),
        "The company observes all gazetted public holidays announced by the Government of Pakistan. Retail staff who work on a public holiday "
        "receive a replacement day off within 30 days.",
        ("h", "6. Remote work"),
        "After confirmation, office staff may work from home up to 2 days per week with their line manager's approval. Remote work is not "
        "available for warehouse and retail roles. Employees working remotely must be reachable on Microsoft Teams during office hours.",
        ("h", "7. Medical benefits"),
        "Confirmed employees receive outpatient (OPD) medical reimbursement of up to PKR 60,000 per year. Hospitalisation is covered by group "
        "health insurance up to PKR 500,000 per year for the employee, spouse and up to two children.",
        ("h", "8. Reimbursements"),
        "Business travel and other work expenses are reimbursed when claims are submitted within 30 days of the expense, with original receipts, "
        "through the Finance portal. Claims older than 30 days are not reimbursed. Approved claims are paid with the next monthly salary.",
        "Employees who use their own car for company work receive a fuel allowance of PKR 45 per kilometre, based on the mileage log approved by "
        "their manager.",
        ("h", "9. Salary and payslips"),
        "Salaries are paid by bank transfer on the 1st working day of each month. Payslips are available on the HR portal. Questions about pay "
        "should be sent to hr@nimbus-appliances.example.",
        ("h", "10. Code of conduct"),
        "Employees must treat colleagues and customers with respect. Harassment of any kind is not tolerated. Concerns can be reported "
        "confidentially to the Harassment Inquiry Committee at conduct@nimbus-appliances.example or to any HR business partner.",
        "Accepting gifts from suppliers worth more than PKR 5,000 is not allowed and must be reported to HR.",
    ])


def it_guide():
    build_pdf("IT_Help_Guide.pdf", "IT Help Guide", "Updated March 2026", [
        ("h", "1. Getting help"),
        "The IT helpdesk is open Monday to Saturday, 9:00 am to 7:00 pm. Raise a ticket by emailing helpdesk@nimbus-appliances.example or by "
        "calling extension 250. Urgent issues (no internet at a branch, POS down) should be reported by phone.",
        ("h", "2. Passwords"),
        "Reset a forgotten password yourself at the self-service portal: https://reset.nimbus-appliances.example. You will need your employee ID "
        "and the phone number registered with HR. Passwords must be at least 12 characters and must be changed every 90 days. The last 5 "
        "passwords cannot be reused. Never share your password, including with IT staff.",
        ("h", "3. VPN access"),
        "To connect to company systems from home, install FortiClient VPN from the Software Center and connect to vpn.nimbus-appliances.example "
        "with your normal login. VPN access is available only to staff approved for remote work.",
        ("h", "4. Wi-Fi"),
        "Staff devices connect to the network 'Nimbus-Staff' using their normal login. Visitors use 'Nimbus-Guest'; the daily guest password is "
        "available at reception.",
        "PAGEBREAK",
        ("h", "5. Laptops and repairs"),
        "If your laptop is slow, broken or damaged, raise a helpdesk ticket. IT aims to respond within 4 working hours and provides a loaner "
        "laptop when a repair takes more than 1 working day. A lost or stolen laptop must be reported to IT within 24 hours so it can be locked "
        "remotely.",
        ("h", "6. Software"),
        "Approved software can be installed from the Software Center without admin rights. Any other software requires a helpdesk ticket and "
        "approval from your line manager.",
        ("h", "7. Email on your phone"),
        "Install Microsoft Outlook from the App Store or Play Store and sign in with your company email. The Microsoft Authenticator app is "
        "required for two-step sign-in.",
        ("h", "8. Printers"),
        "Head office printers are on every floor. Send jobs to 'Nimbus-Print' and release them at any printer with your ID card. Colour printing "
        "requires manager approval.",
        ("h", "9. Security"),
        "Forward suspicious emails to phishing@nimbus-appliances.example and then delete them. Do not click links or open attachments you did not "
        "expect. USB storage devices are blocked on company laptops.",
    ])


def catalogue():
    rows = [["Model", "Product", "Price (PKR)", "Warranty"],
            ["NA-150i", "Nimbus AirCool 1.5 Ton Inverter AC", "189,000", "10 years compressor, 2 years parts"],
            ["NA-100i", "Nimbus AirCool 1 Ton Inverter AC", "154,500", "10 years compressor, 2 years parts"],
            ["NF-18", "Nimbus FrostFree Refrigerator 18 cu ft", "139,500", "10 years compressor, 1 year parts"],
            ["NF-14", "Nimbus FrostFree Refrigerator 14 cu ft", "112,900", "10 years compressor, 1 year parts"],
            ["NW-8F", "Nimbus WashPro 8 kg Front Load Washer", "112,000", "5 years motor, 1 year parts"],
            ["NW-10T", "Nimbus WashPro 10 kg Top Load Washer", "74,500", "5 years motor, 1 year parts"],
            ["NM-25", "Nimbus Microwave Oven 25 L", "29,900", "1 year"],
            ["NAF-55", "Nimbus Air Fryer 5.5 L", "24,999", "1 year"],
            ["NH-2000", "Nimbus Ceramic Room Heater 2000 W", "18,500", "1 year"],
            ["NWD-3", "Nimbus Water Dispenser 3 Taps", "46,000", "1 year, 3 years compressor"]]
    build_pdf("Product_Catalogue.pdf", "Product Catalogue 2026", "Prices valid from 1 September 2026", [
        ("h", "Our range"),
        "Nimbus designs energy-efficient home appliances for Pakistani homes. All prices include sales tax. Prices and availability may change; "
        "the price on the invoice at the time of purchase applies.",
        ("table", rows, [22 * mm, 68 * mm, 25 * mm, 55 * mm]),
        ("h", "Air conditioners"),
        "The AirCool Inverter series uses a DC inverter compressor that saves up to 60% electricity compared with non-inverter models. Both "
        "models run on low voltage down to 160 V, include a Wi-Fi remote app and use R32 refrigerant. Standard installation (up to 10 feet of "
        "copper pipe) is free in Lahore, Karachi and Islamabad. Extra pipe is charged at PKR 1,200 per foot.",
        "PAGEBREAK",
        ("h", "Refrigerators"),
        "FrostFree refrigerators have an inverter compressor, a 5-in-1 convertible freezer and a 10-year compressor warranty. The 18 cu ft model "
        "has an A+ energy rating and a door alarm.",
        ("h", "Washing machines"),
        "The WashPro 8 kg front load washer offers 15 programs including Quick 15-minute wash and steam hygiene. The 10 kg top load model is "
        "designed for large families and has a child lock.",
        ("h", "Small appliances"),
        "The Air Fryer 5.5 L cooks with up to 85% less oil and has 8 presets. The Microwave Oven 25 L has grill and defrost functions. The Ceramic "
        "Room Heater has tip-over protection and 3 heat settings.",
        ("h", "Installments"),
        "All products above PKR 50,000 can be bought on 0% installments for 12 months through partner banks (credit cards of HBL, Meezan Bank "
        "and Bank Alfalah).",
    ])


def refund_policy():
    build_pdf("Refund_and_Warranty_Policy.pdf", "Refund, Return and Warranty Policy", "Version 2.0, effective 1 July 2026", [
        ("h", "1. Returns"),
        "Unused products in their original, unopened packaging can be returned within 7 days of delivery for a full refund. Installed "
        "appliances (air conditioners, water dispensers) cannot be returned unless they are faulty.",
        ("h", "2. Damaged or wrong items"),
        "If a product arrives damaged or you receive the wrong item, report it within 48 hours of delivery by calling our helpline or emailing "
        "support@nimbus-appliances.example with photos of the product and packaging. We will collect it and send a replacement free of charge.",
        ("h", "3. Refunds"),
        "Approved refunds are processed within 10 working days to the original payment method. Card refunds may take a further 5-7 days to "
        "appear on your statement, depending on your bank. Cash on delivery (COD) orders are refunded by bank transfer to an account in the "
        "customer's name.",
        "Delivery charges are refunded only if the product was damaged, faulty or wrong.",
        "PAGEBREAK",
        ("h", "4. Warranty"),
        "Warranty periods are listed in the Product Catalogue for each model and start on the date of delivery. To make a warranty claim, contact "
        "the helpline with your invoice number and the product's serial number (printed on the label at the back or side of the product). A "
        "technician visit is arranged within 72 hours in major cities.",
        "The warranty does not cover damage from voltage fluctuation without a stabiliser (except inverter models rated for low voltage), "
        "physical damage, or repairs by anyone other than a Nimbus-authorised technician. Unauthorised repairs make the warranty void.",
        ("h", "5. Exchanges"),
        "Faulty products that cannot be repaired within 15 days are exchanged for the same model, or a model of equal value if it is out of stock.",
    ])


def shipping_docx():
    from docx import Document
    d = Document()
    d.core_properties.title = "Shipping and Delivery Policy"
    d.add_heading("Shipping and Delivery Policy", 0)
    d.add_paragraph(f"{COMPANY} · Updated August 2026")
    sections = [
        ("Delivery times", "Orders are delivered within 2-4 working days in Lahore, Karachi, Islamabad and Rawalpindi, and within 5-7 working "
                           "days to other cities. Orders placed after 3:00 pm are processed the next working day."),
        ("Delivery charges", "Delivery is free for orders above PKR 20,000. A flat delivery charge of PKR 500 applies to smaller orders."),
        ("Cash on delivery", "Cash on delivery (COD) is available for orders up to PKR 150,000. Larger orders must be paid in advance by card, "
                             "bank transfer, JazzCash or Easypaisa."),
        ("Tracking", "You receive an SMS with a tracking link when your order is dispatched, and a call from the rider before delivery."),
        ("Installation", "Air conditioners and water dispensers are installed by our technicians. Installation is scheduled within 48 hours of "
                         "delivery; the technician calls you to fix a time."),
        ("Failed delivery", "If nobody is available, the rider tries again the next working day. After two failed attempts the order returns "
                            "to our warehouse and COD orders are cancelled."),
    ]
    for h, text in sections:
        d.add_heading(h, level=1)
        d.add_paragraph(text)
    d.add_paragraph(SAMPLE_NOTE)
    d.save(OUT / "Shipping_and_Delivery_Policy.docx")


def faq_txt():
    (OUT / "Company_FAQ.txt").write_text(f"""Company FAQ - {COMPANY}
{SAMPLE_NOTE}

Where is your head office?
Our head office is at 42-C Gulberg III, Lahore. We also have experience centres in Karachi (Shahrah-e-Faisal) and Islamabad (Blue Area).

What are your customer support hours?
Customer support is available Monday to Saturday, 9:00 am to 8:00 pm, on the helpline 0800-64662 (toll free) and by email at support@nimbus-appliances.example.

Which payment methods do you accept?
We accept debit and credit cards, bank transfer, JazzCash, Easypaisa and cash on delivery. Products above PKR 50,000 are available on 0% installments for 12 months with partner banks.

Do you offer corporate or bulk discounts?
Yes. Orders of 10 or more units qualify for corporate pricing. Email corporate@nimbus-appliances.example for a quotation.

How do I apply for a job at Nimbus?
Current openings are posted on our careers page and on LinkedIn. You can also send your CV to careers@nimbus-appliances.example.

Do you have a service centre?
Yes, authorised service centres operate in Lahore, Karachi, Islamabad, Faisalabad and Multan. Book a technician through the helpline.
""")


def salary_bands():
    build_pdf("Salary_Bands_2026.pdf", "Salary Bands 2026 (Confidential)", "Management only - do not share", [
        ("h", "Monthly gross salary bands"),
        ("table", [["Grade", "Role examples", "Minimum (PKR)", "Maximum (PKR)"],
                   ["G1", "Sales associate, warehouse assistant", "45,000", "70,000"],
                   ["G2", "Senior associate, technician", "70,000", "110,000"],
                   ["G3", "Officer, team lead", "110,000", "180,000"],
                   ["G4", "Assistant manager", "180,000", "280,000"],
                   ["G5", "Manager", "280,000", "450,000"]], [18 * mm, 72 * mm, 35 * mm, 35 * mm]),
        "Annual increments are between 8% and 15% based on the performance rating. This document is restricted to the management team.",
    ])


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    for f in OUT.iterdir():
        f.unlink()
    hr_policy(); it_guide(); catalogue(); refund_policy(); shipping_docx(); faq_txt(); salary_bands()
    print("Wrote", sorted(p.name for p in OUT.iterdir()))


if __name__ == "__main__":
    main()
